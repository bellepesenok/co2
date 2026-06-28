# -*- coding: utf-8 -*-
# Generate exam cheat-sheet (agency law) as a Word document.

from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement


CN_FONT = '宋体'
CN_FONT_H = '黑体'

doc = Document()

normal = doc.styles['Normal']
normal.font.name = CN_FONT
normal.font.size = Pt(10.5)
normal.element.rPr.rFonts.set(qn('w:eastAsia'), CN_FONT)


def set_cn(run, font=CN_FONT):
    run.font.name = font
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.find(qn('w:rFonts'))
    if rFonts is None:
        rFonts = OxmlElement('w:rFonts')
        rPr.append(rFonts)
    rFonts.set(qn('w:eastAsia'), font)
    rFonts.set(qn('w:ascii'), font)
    rFonts.set(qn('w:hAnsi'), font)


def shade_cell(cell, color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), color)
    tcPr.append(shd)


RED = RGBColor(0xC0, 0x39, 0x2B)
BLUE = RGBColor(0x1F, 0x3A, 0x6E)
GREEN = RGBColor(0x1E, 0x7D, 0x32)


def title(text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(text)
    set_cn(r, CN_FONT_H)
    r.font.size = Pt(20)
    r.font.bold = True
    r.font.color.rgb = BLUE


def subtitle(text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(text)
    set_cn(r)
    r.font.size = Pt(11)
    r.font.color.rgb = RGBColor(0x66, 0x66, 0x66)


def h1(text):
    p = doc.add_paragraph()
    r = p.add_run(text)
    set_cn(r, CN_FONT_H)
    r.font.size = Pt(15)
    r.font.bold = True
    r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    pPr = p._p.get_or_add_pPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:fill'), '1F3A6E')
    pPr.append(shd)
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(6)


def h2(text):
    p = doc.add_paragraph()
    r = p.add_run(text)
    set_cn(r, CN_FONT_H)
    r.font.size = Pt(12.5)
    r.font.bold = True
    r.font.color.rgb = RED
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(3)


def para(runs, bullet=False):
    p = doc.add_paragraph(style='List Bullet' if bullet else None)
    if isinstance(runs, str):
        runs = [(runs, False, None)]
    for item in runs:
        if isinstance(item, str):
            text, bold, color = item, False, None
        else:
            text, bold, color = (list(item) + [False, None])[:3]
        r = p.add_run(text)
        set_cn(r)
        r.font.size = Pt(10.5)
        r.font.bold = bold
        if color:
            r.font.color.rgb = color
    p.paragraph_format.space_after = Pt(2)


def make_table(headers, rows, widths=None, header_fill='1F3A6E'):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = 'Table Grid'
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr = table.rows[0].cells
    for i, htext in enumerate(headers):
        hdr[i].text = ''
        rr = hdr[i].paragraphs[0].add_run(htext)
        set_cn(rr, CN_FONT_H)
        rr.font.size = Pt(10)
        rr.font.bold = True
        rr.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        shade_cell(hdr[i], header_fill)
    for ri, row in enumerate(rows):
        cells = table.add_row().cells
        fill = 'EEF2F8' if ri % 2 == 0 else 'FFFFFF'
        for i, ctext in enumerate(row):
            cells[i].text = ''
            for j, line in enumerate(str(ctext).split('\n')):
                pp = cells[i].paragraphs[0] if j == 0 else cells[i].add_paragraph()
                rr = pp.add_run(line)
                set_cn(rr)
                rr.font.size = Pt(9.5)
            shade_cell(cells[i], fill)
    if widths:
        for i, w in enumerate(widths):
            for row in table.rows:
                row.cells[i].width = Cm(w)
    return table


LQ = '\u300c'  # 「
RQ = '\u300d'  # 」


# ========================== 封面 ==========================
title('代理法 · 考试速记 & 答题模板')
subtitle('代理类型判断 · 是否构成代理 · 无权代理与表见代理 · 直接/间接代理 · 案例标准答案')
subtitle('（依据《中华人民共和国民法典》及英美法、大陆法代理制度整理）')

# ========================== 一、核心概念速记 ==========================
h1('一、核心概念速记（先记牢这几句）')

h2('1. 代理的定义（核心是' + LQ + '授权' + RQ + '）')
para([('代理', True, RED), ('：代理人（agent）按照本人（principal）的', False, None),
      ('授权（authorization）', True, BLUE),
      ('，以本人名义同第三人订立合同或作其他法律行为，所产生的权利义务', False, None),
      ('直接对本人发生效力', True, RED), ('。', False, None)])
para([('三方结构', True, RED), ('：被代理人（本人）—— 代理人 —— 相对人（第三人），三者缺一不可。', False, None)])
para([('核心要件', True, BLUE), ('：①有代理权（授权）；②以本人名义；③在权限范围内；④实施法律行为。', False, None)])

h2('2. 三个最易混淆的概念（必考辨析）')
make_table(
    ['概念', '是否独立作出意思表示', '是否需行为能力', '法律结构', '举例'],
    [
        ['代理人', '独立决定意思表示内容\n（买什么、多少钱由其定）', '须有相应行为能力\n无行为能力人不得当代理人', '三方结构', '律师受托在500万内自行选购文物'],
        ['使者（传达人）', '不能独立，只传达本人\n已决定的意思', '不需要行为能力\n5岁儿童也可当使者', '本质仍是本人与第三人', '父亲让5岁儿子去转达' + LQ + '以435万买瓷瓶' + RQ],
        ['冒名行为', '冒用他人姓名为自己谋利', '—', '双方结构\n（冒名人与相对人）', '丙谎称自己是甲向乙借款'],
    ],
    widths=[2.3, 3.2, 3.0, 2.4, 3.6],
)
para([('记忆口诀：', True, RED), ('使者' + LQ + '传话' + RQ + '、代理' + LQ + '做主' + RQ + '、冒名' + LQ + '顶替' + RQ + '。冒名只有两方，不是代理！', False, None)])

# ========================== 二、判断流程 ==========================
h1('二、判断流程图：是不是代理？属于哪种代理？')

h2('第一步：是否构成代理？（三连问）')
para([('① 是否有三方结构（本人、代理人、第三人）？', True, BLUE)])
para('　 否 → 冒名行为（双方结构），不是代理。', bullet=True)
para([('② 行为人是否以' + LQ + '本人名义' + RQ + '行事？', True, BLUE)])
para('　 以自己名义 → 可能是间接代理/行纪（见第五部分），不是直接代理。', bullet=True)
para([('③ 是否有代理权（授权、法定或被追认）？', True, BLUE)])
para('　 无代理权 → 进入' + LQ + '无权代理 / 表见代理' + RQ + '判断（第四部分）。', bullet=True)

h2('第二步：判断代理' + LQ + '权来源' + RQ + '的类型')
make_table(
    ['法系', '类型', '判断标准（看到什么就选它）'],
    [
        ['大陆法', '法定代理', '依法律规定（如监护人）、法院选任（破产清算人）、私人选任（监护人、遗产管理人）取得代理权'],
        ['大陆法', '意定代理（委托代理）', '由本人意思表示授予；口头或书面均可；可向代理人或第三人表示'],
        ['英美法', '明示授权 express', '本人口头或书面明确授权（最典型）'],
        ['英美法', '默示授权 implied', '无明确授权，但本人常以言行让代理人以本人名义活动并担责（习惯性、默许）'],
        ['英美法', '不容否认的代理\nestoppel（=表见代理）', '代理人无授权，但其行为足以令第三人确信有代理权，且本人知情而不否认'],
        ['英美法', '客观必需的代理\nof necessity', '紧急状态下，为维护本人利益而不得不代理（须本人利益+紧急+无法联系本人）'],
        ['英美法', '追认的代理\nratification', '对原本无权的代理，本人事后追认使其生效'],
    ],
    widths=[1.6, 3.4, 9.5],
)

# ========================== 三、英美法类型判别 ==========================
h1('三、英美法代理类型——见招拆招（PPT练习答案）')
make_table(
    ['情境关键词', '代理类型', '理由'],
    [
        ['签合同、明确委托', '明示授权', '皮特书面委托克鲁斯赴华买奶粉'],
        ['习惯让…替他做、常亲自对账', '默示授权', '科比习惯让姚明销售球衣并亲自清账'],
        ['一直知情、长期默许他人以我名义', '不容否认/表见代理', '世贤一直知情品如向艾莉出售其设计方案'],
        ['暴雨/急病/突发、为本人利益、联系不上', '客观必需的代理', 'Angela Baby暴雨中紧急组织搬运演出设备'],
        ['事先无权，事后是否承认', '追认的代理 / 无权代理', '小S以大S名义办宴请，看大S是否追认'],
    ],
    widths=[4.5, 4.0, 6.0],
)

# ========================== 四、无权代理 vs 表见代理 ==========================
h1('四、无权代理 vs 表见代理（核心考点·必背对比表）')
make_table(
    ['对比项', '狭义无权代理', '表见代理（不容否认的代理）'],
    [
        ['概念', '欠缺代理权而实施的代理行为：自始无权、超越权限、代理权终止后', '无权代理人以本人名义订约，但相对人有理由相信其有代理权'],
        ['构成要件',
         '①行为人欠缺代理权\n②除欠权外无其他无效事由\n（若违反效力性强制规定/公序良俗→无效代理）',
         '①无代理权\n②有权利外观（介绍信、盖章空白合同等）\n③相对人善意且无过失\n④权利外观的形成可归责于被代理人'],
        ['效力', '效力待定：经本人追认才对本人生效；未追认则对本人不生效', '代理行为有效，直接对被代理人发生效力'],
        ['相对人权利', '可催告本人（30日内）追认；善意相对人在追认前有撤销权', '可主张有权代理效果；亦可在追认生效前行使撤销权（有选择权）'],
        ['立法目的', '侧重保护被代理人利益', '侧重保护善意相对人利益与交易安全'],
        ['法条', '民法典第171条（原合同法48条）', '民法典第172条（原合同法49条）'],
    ],
    widths=[2.2, 6.0, 6.3],
)

h2('★ 表见代理不成立的情形（权利外观不可归责于被代理人）')
para('① 行为人伪造他人公章、合同书、授权委托书，假冒名义实施——如' + LQ + '老干妈案' + RQ + '。', bullet=True)
para('② 公章/合同书/授权书遗失、被盗，或职务关系已终止并已合理公告或通知、相对人应当知悉。', bullet=True)

h2('★ 无权代理的英美法处理（违反' + LQ + '有代理权的默示担保' + RQ + '）')
para('代理人对第三人有' + LQ + '保证自己有代理权' + RQ + '的默示担保；无权代理即违反该担保。', bullet=True)
para([('严格责任', True, RED), ('：不论代理人主观如何均需负责。', False, None)], bullet=True)
para('免责情形：第三人知道其无权/未提供担保；合同已排除代理人责任；本人指示含糊而代理人善意合理执行。', bullet=True)
para([('赔偿范围', True, BLUE), ('：不超过' + LQ + '假设合同因追认而生效后第三人所能获得的履行利益' + RQ + '。', False, None)], bullet=True)

# ========================== 五、直接代理 vs 间接代理 ==========================
h1('五、本人/代理人与第三人的关系——直接代理 vs 间接代理（外部关系）')

h2('1. 大陆法')
make_table(
    ['类型', '以谁名义', '合同效力归属'],
    [
        ['直接代理', '以被代理人名义', '合同效力直接归于被代理人'],
        ['间接代理', '以代理人自己名义', '代理人承担权利义务；代理人可转让，则被代理人取代其地位'],
    ],
    widths=[3.0, 4.0, 7.5],
)

h2('2. 英美法（按' + LQ + '义务标准' + RQ + '——谁对合同负责）')
make_table(
    ['订约时披露情况', '英文', '效果'],
    [
        ['指明了本人姓名', 'agent for a named principal', '效果归于本人'],
        ['表示有代理关系但未指明本人', 'agent for an unnamed principal', '效果仍归属于本人'],
        ['完全不披露代理关系（隐名本人）', 'undisclosed principal',
         '①代理人负责；②未披露的本人原则上可介入，直接取得合同权利义务：\n　a.本人可直接向第三人请求或起诉；\n　b.第三人知道本人存在后有选择权，可要求本人或代理人负责，一经选定不得变更。\n③本人介入权的限制：介入与合同条款相抵触、或第三人基于对代理人的信赖才订约的，不得介入'],
    ],
    widths=[4.0, 4.3, 7.0],
)

h2('3. 中国外贸代理制（隐名代理）——民法典第925、926条')
para([('第925条（原合同法402）：', True, BLUE),
      ('受托人以自己名义、在授权范围内与第三人订约，第三人订约时', False, None),
      ('知道', True, RED), ('受托人与委托人代理关系的→合同', False, None),
      ('直接约束委托人和第三人', True, RED),
      ('；但有确切证据证明只约束受托人与第三人的除外。', False, None)])
para([('第926条（原合同法403）：', True, BLUE),
      ('受托人以自己名义订约，第三人', False, None), ('不知道', True, RED),
      ('代理关系的：', False, None)])
para('　A. 受托人因第三人原因对委托人不履行→应向委托人披露第三人，委托人可行使受托人对第三人的权利（介入权）；但第三人若知委托人就不订约的除外。', bullet=True)
para('　B. 受托人因委托人原因对第三人不履行→应向第三人披露委托人，第三人可选择受托人或委托人为相对人主张权利（选择权），但不得变更选定的相对人。', bullet=True)

# ========================== 六、终止与义务 ==========================
h1('六、代理关系的终止 & 代理人/本人义务')

h2('1. 代理关系终止')
para([('依意思表示终止：', True, BLUE), ('①约定期限届满；②双方同意；③本人单方撤回（须提前通知；若授权与代理人利益结合则不得单方撤回）。', False, None)])
para([('依法律规定终止：', True, BLUE), ('本人或代理人死亡、破产、丧失行为能力（商事代理有例外）。', False, None)])
para([('终止的效果：', True, BLUE), ('对内——代理人失去代理权、应给予适当补偿；对外——', False, None),
      ('须通知第三人才生效', True, RED), ('，未通知不得对抗第三人的履约请求，本人仍须负责，但可向代理人追偿。', False, None)])

h2('2. 代理人义务')
para('①勤勉履行（duty of care）；②诚信忠实（good faith & loyalty）：不得自己代理/双方代理（除非本人同意或追认）、须公开客户必要情况、不得受贿或谋私、不得串通损害本人。', bullet=True)
para('③保密义务；④申报账目；⑤原则上不得复代理（转委托）。', bullet=True)
para([('复代理三例外：', True, RED), ('①本人事先同意；②本人事后追认；③紧急情况下为维护本人利益必须转委托（急病、通讯中断等）。', False, None)], bullet=True)

h2('3. 本人义务')
para('①支付佣金/报酬；②偿还代理人履行代理产生的费用、赔偿其损失；③让代理人查核账册。', bullet=True)

# ========================== 七、答题模板 ==========================
h1('七、考试答题模板（套用即可）')

h2('模板A：判断' + LQ + '是否构成代理 / 是否为代理人' + RQ)
para('第一步【定性】：本案中，X以Y的名义、就……事项与Z实施法律行为，涉及代理之认定。')
para('第二步【三要件】：判断是否符合代理三方结构（本人/代理人/相对人）、是否以本人名义、是否在代理权限内。')
para('第三步【辨析】：区分代理与使者（是否独立作出意思表示）、与冒名行为（是否三方结构）、与居间/行纪（以谁名义、是否撮合）。')
para([('第四步【结论】：', True, RED), ('故X（是/否）Y的代理人，本行为（构成/不构成）代理。', False, None)])

h2('模板B：判断代理权类型')
para('先定法系（题目语境）→ 大陆法分法定/意定；英美法对照' + LQ + '明示、默示、表见、客观必需、追认' + RQ + '五型 → 指出关键事实（如' + LQ + '一直知情' + RQ + '、' + LQ + '暴雨紧急' + RQ + '、' + LQ + '事后追认' + RQ + '）→ 得出类型。')

h2('模板C：无权代理 / 表见代理分析')
para('①确认行为人无代理权（自始无权/超越/终止后）。')
para('②是否构成表见代理：逐一检验——有无权利外观？相对人是否善意无过失？权利外观能否归责于被代理人？')
para([('③下结论：', True, RED),
      ('若四要件齐备→构成表见代理，代理行为有效，被代理人担责后可向无权代理人追偿；', False, None)])
para('　若权利外观不可归责（如伪造公章）→不构成表见代理，属狭义无权代理，效力待定，未经追认对本人不生效；善意相对人可向无权代理人主张责任（英美法：违反默示担保，赔偿以履行利益为限）。')

h2('模板D：间接代理（隐名/未披露本人）分析')
para('①认定受托人以自己名义、在授权范围内与第三人订约。')
para('②看第三人订约时是否知道代理关系：')
para('　知道→民法典925条，合同直接约束委托人与第三人（除非有确切证据只约束受托人与第三人）。', bullet=True)
para('　不知道→合同相对性原则先约束受托人与第三人；再依926条，因第三人原因不履行→委托人介入权；因委托人原因不履行→第三人选择权（选定后不得变更）。', bullet=True)
para([('③结论。', True, RED)])

# ========================== 八、案例标准答案 ==========================
h1('八、PPT案例·标准答案速查')

cases = [
    ('【选择题】下列哪些属于代理？（甲请乙代购等）',
     'A 乙以甲名义代购→就甲的部分构成代理；B 乙把纸条交给销售员并称为朋友买→乙是使者（传达），非代理；D 介绍歌星签三方协议→居间/介绍，非代理。结论：A构成代理。'),
    ('【外运公司代办焦炭出口】付费是否合理？',
     '不合理。货代将货装船即已完成代理职责，承运人未收运费应向货主（提单托运人）主张；外运公司与承运人无合约关系、无付款义务。其付费源于对自身法律地位（货代权义）不清楚。'),
    ('【E公司/R公司·皮货案】R是否有代理权？',
     '国际商事代理权类型：明示、默示、表见、客观必需、追认。皮货并非易腐或久存大损价值之物，R越权高价出售不能认定为' + LQ + '为本人最大利益' + RQ + '，不具备客观必需的代理权，应对越权造成E公司的损失负责。'),
    ('【选择题】哪一情形构成无权代理？',
     'A 甲冒用乙姓名领稿酬→冒名行为；C 刘某冒充丁某面试→冒名行为（且非财产法律行为）；D 关某以邻居李某名义代收并代付保健品→以他人名义实施、三方结构→构成无权代理。结论：D。'),
    ('【乙长期签单餐费·辞职后被拒付】',
     '乙作为业务经理长期签单、公司按月结清，已形成权利外观且可归责于公司→构成表见代理，甲公司应当付款（不能拒付）。公司付款后可向乙追偿；相对人不能直接要求乙承担连带或补充责任。'),
    ('【嫣然公司·亚鹏菲姐签单案】是否还钱？',
     '应还。亚鹏系中层管理人员长期以公司名义签单消费、公司确认菲姐部分→对亚鹏部分亦形成可归责的权利外观，构成表见代理（不容否认的代理），嫣然公司应偿付。'),
    ('【肖玉芬/肖德强·委托炒股案】',
     '肖德强未获其父授权，但基于父子关系及其将股票资料告知贾静、支付佣金等行为，足以使贾静相信其有代理权→表见代理成立，委托资产管理协议有效，肖玉芬应依约支付管理费。'),
    ('【老干妈案】是否成立代理？协议是否有效？是否付费？',
     '①不成立有效代理：曹某等伪造公章、冒充经理（甚至涉刑），权利外观（伪造公章）不可归责于老干妈，不构成表见代理，属狭义无权代理/冒名。②协议未经老干妈追认，对其不发生效力。③老干妈无需付费。'),
    ('【新加坡甲/泰国丙伪造公章·英美法】',
     '①伪造公章不可归责于甲→不构成表见代理，属狭义无权代理；甲拒绝追认，乙不得请求甲履行。②乙为善意相对人，可以丙违反' + LQ + '有代理权的默示担保' + RQ + '起诉丙，丙须担责。③赔偿范围不超过合同经追认生效后乙的履行利益。'),
    ('【A/B外运公司·货物淋湿案】谁赔偿？',
     'B外运公司受A委托，负有代理人的谨慎义务（妥善照料、对露天货物必要苫盖）。其未妥善苫盖致货物淋湿残损，违反duty of care，应承担淋湿赔偿责任。'),
    ('【甲委托乙卖房·乙自己买】乙该不该退房？',
     '乙构成自己代理（代理本人与自己交易），属滥用代理权，除本人事前同意或事后追认外无效。甲坚持退房不予追认→乙的行为为无效代理，应退房。若甲不要求退房（即追认）→行为有效。'),
    ('【E/R公司·生鲜+选任C公司·复代理】C是否有代理权？',
     '原则上代理人无复任权；例外：①本人事先同意；②事后追认；③紧急情况下为维护本人利益必须转委托。本案生鲜将腐、无法联系E公司，属紧急必要情形，R选任C构成有效转委托，C具备代理权。'),
    ('【Vacheron手表案·丙知情】',
     '乙在权限内以自己名义与知情的丙订约（为甲计算）→适用民法典925条，甲' + LQ + '自动取代' + RQ + '，买卖合同直接约束甲、丙；除非有确切证据证明只约束乙、丙（如约定' + LQ + '丙只认乙' + RQ + '）。'),
    ('【Vacheron手表案·丙不知情】',
     '依合同相对性，合同先约束乙、丙。大陆法：乙可将合同权益转让给甲，甲取代乙地位；英美法：未披露的本人甲可介入，直接取得权利、向丙请求交付。若乙因甲原因不付款→乙须向丙披露甲，丙享有选择权（选定后不得变更）。'),
    ('【选择题·甲委托乙以乙名义购机械设备，丙不知情，乙因甲原因未付款】',
     '乙以自己名义订约、丙不知代理关系、乙因委托人（甲）原因不履行→依926条乙应披露委托人，丙可选择甲或乙为相对人主张权利。结论：C（可选择要求甲或乙支付）。'),
    ('【职务行为案例·王某 / 赵某】',
     '案例一：王某以自己名义出具欠款条、未告知用途、未让对方开发票，李某不知是职务行为→对外类似隐名/个人交易，王某应承担还款责任（再向单位追偿）。案例二：赵某虽以自己名义出具欠款条，但电器公司将空调安装在酒店、向酒店开具发票且酒店已入账→已披露并实际归属单位，构成职务行为，由酒店（用人单位）担责。'),
]

for q, a in cases:
    p = doc.add_paragraph()
    r = p.add_run(q)
    set_cn(r, CN_FONT_H)
    r.font.size = Pt(10.5)
    r.font.bold = True
    r.font.color.rgb = BLUE
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(1)
    pa = doc.add_paragraph()
    ra = pa.add_run('答：')
    set_cn(ra)
    ra.font.bold = True
    ra.font.color.rgb = RED
    ra.font.size = Pt(10.5)
    ra2 = pa.add_run(a)
    set_cn(ra2)
    ra2.font.size = Pt(10.5)
    pa.paragraph_format.space_after = Pt(4)

# ========================== 九、关键法条 ==========================
h1('九、关键法条索引（背诵备用）')
make_table(
    ['条文', '内容要点'],
    [
        ['民法典161-163', '可通过代理实施民事法律行为；代理人在权限内以本人名义实施→对本人生效；代理含委托代理与法定代理'],
        ['民法典168', '禁止自己代理、双方代理，除非本人同意或追认'],
        ['民法典169', '转委托（复代理）须本人同意或追认'],
        ['民法典171', '无权代理：未经追认对本人不生效；相对人催告权、善意相对人撤销权'],
        ['民法典172', '表见代理：相对人有理由相信有代理权→代理行为有效'],
        ['民法典173-175', '委托代理、法定代理的终止情形'],
        ['民法典925（原合同法402）', '隐名代理·第三人知情→合同直接约束委托人与第三人'],
        ['民法典926（原合同法403）', '未披露本人·第三人不知情→委托人介入权 / 第三人选择权'],
    ],
    widths=[4.5, 10.5],
)

doc.save('/workspace/法律考试资料/代理法_考试速记与答题模板.docx')
print('OK saved')
