# -*- coding: utf-8 -*-
"""
краткая упрощённая версия конспекта главы 6 для иностранного студента.
简明版：简单中文 + 俄语 + 拼音，方便外国学生背重点。
"""
from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

CN_FONT = "宋体"
LAT_FONT = "Times New Roman"

C_TITLE = RGBColor(0x1F, 0x38, 0x64)
C_H1 = RGBColor(0xC0, 0x00, 0x00)
C_CN = RGBColor(0x00, 0x00, 0x00)
C_PY = RGBColor(0x80, 0x80, 0x80)
C_RU = RGBColor(0x1F, 0x4E, 0x2A)
C_KEY = RGBColor(0xB0, 0x00, 0x00)

doc = Document()
style = doc.styles["Normal"]
style.font.name = LAT_FONT
style.font.size = Pt(11)
style._element.rPr.rFonts.set(qn("w:eastAsia"), CN_FONT)


def _set(run, font, size, bold=None, italic=None, color=None):
    run.font.name = font
    run._element.rPr.rFonts.set(qn("w:eastAsia"), font)
    run.font.size = Pt(size)
    if bold is not None:
        run.font.bold = bold
    if italic is not None:
        run.font.italic = italic
    if color is not None:
        run.font.color.rgb = color


def shade(p, hexc):
    pPr = p._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), hexc)
    pPr.append(shd)


def title(text, sub):
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _set(p.add_run(text), CN_FONT, 22, bold=True, color=C_TITLE)
    p2 = doc.add_paragraph(); p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _set(p2.add_run(sub), LAT_FONT, 12, italic=True, color=RGBColor(0x55,0x55,0x55))


def h1(text):
    p = doc.add_paragraph(); p.space_before = Pt(12)
    _set(p.add_run(text), CN_FONT, 15, bold=True, color=C_H1)
    shade(p, "F2DCDB")


def point(cn, py, ru):
    """один пункт: китайский (жирный) + пиньинь (серый) + русский (новая строка)."""
    p = doc.add_paragraph(style="List Bullet")
    _set(p.add_run(cn), CN_FONT, 12, bold=True, color=C_CN)
    if py:
        _set(p.add_run("  " + py), LAT_FONT, 10, italic=True, color=C_PY)
    pr = doc.add_paragraph()
    pr.paragraph_format.left_indent = Cm(0.9)
    _set(pr.add_run("→ " + ru), LAT_FONT, 10.5, color=C_RU)


def note(ru, header="Запомни / 记住"):
    p = doc.add_paragraph()
    _set(p.add_run("✓ " + header + ": "), LAT_FONT, 10.5, bold=True, color=RGBColor(0x8a,0x60,0x00))
    _set(p.add_run(ru), LAT_FONT, 10.5, color=RGBColor(0x5a,0x42,0x00))
    shade(p, "FFFBEA")


def make_table(headers, rows, widths):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Light Grid Accent 1"; t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, htext in enumerate(headers):
        t.rows[0].cells[i].text = ""
        _set(t.rows[0].cells[i].paragraphs[0].add_run(htext), CN_FONT, 10.5, bold=True, color=RGBColor(0xFF,0xFF,0xFF))
    for row in rows:
        cells = t.add_row().cells
        for i, val in enumerate(row):
            cells[i].text = ""
            _set(cells[i].paragraphs[0].add_run(val), CN_FONT, 10)
    for r in t.rows:
        for i, w in enumerate(widths):
            r.cells[i].width = Cm(w)


# ===================== ТИТУЛ =====================
title("第六章 简明版", "Глава 6. Кратко и просто — для иностранного студента (закрытый экзамен)")
p = doc.add_paragraph()
_set(p.add_run("对外贸易行政管理"), CN_FONT, 14, bold=True, color=C_KEY)
_set(p.add_run("  duìwài màoyì xíngzhèng guǎnlǐ  —  административное управление внешней торговлей"),
     LAT_FONT, 10.5, italic=True, color=RGBColor(0x55,0x55,0x55))
note("Главная мысль: 经济手段为主，行政手段为辅 — основа = экономические методы, "
     "административные методы только В ДОПОЛНЕНИЕ.", header="Самое главное")

# ===================== 1 =====================
h1("1. 概述 — Общее (что это и зачем)")
point("行政管理有4个特点", "yǒu 4 ge tèdiǎn", "У админуправления 4 признака:")
point("统一、速效、强制、纵向", "tǒngyī, sùxiào, qiángzhì, zòngxiàng",
      "единое, быстрое, принудительное, вертикальное (сверху вниз).")
point("管理3个对象", "guǎnlǐ 3 ge duìxiàng", "Управляет 3 объектами:")
point("经营者、货物、配套环节", "jīngyíngzhě, huòwù, pèitào huánjié",
      "субъекты ВЭД, сами товары, сопутствующие звенья.")
point("8种手段", "8 zhǒng shǒuduàn",
      "8 методов: регистрация, госторговля, лицензия, квота, инспекция товаров, происхождение, таможня, валюта.")
point("市场有3个缺点", "shìchǎng yǒu 3 ge quēdiǎn", "У рынка 3 недостатка:")
point("自发性、盲目性、滞后性", "zìfāxìng, mángmùxìng, zhìhòuxìng",
      "стихийность, слепота, запаздывание — поэтому нужно государство.")
note("Тест часто спрашивает: «4 признака» и «3 недостатка рынка». Учи наизусть.")

# ===================== 2 =====================
h1("2. 经营管理 — Кто может торговать")
point("经营权：许可制 → 备案登记制", "xǔkězhì → bèi'àn dēngjìzhì",
      "Право на ВЭД: раньше нужно было РАЗРЕШЕНИЕ, теперь только РЕГИСТРАЦИЯ (уведомление).")
point("不登记，海关不放行", "bù dēngjì, hǎiguān bù fàngxíng",
      "Без регистрации таможня не выпустит товар.")
point("国营贸易 = 专营权", "guóyíng màoyì = zhuānyíngquán",
      "Госторговля = исключительное право: только избранные фирмы торгуют важными товарами.")
point("用目录管理", "yòng mùlù guǎnlǐ", "Управляют через 2 списка: список фирм + список товаров.")
point("进口8种国营货物", "jìnkǒu 8 zhǒng",
      "8 импортных гостоваров: зерно, растит. масло, сахар, табак, нефть, нефтепродукты, удобрения, хлопок.")
note("Главное здесь: 许可制 → 备案登记制 (от разрешения к регистрации) и 8 импортных товаров.")

# ===================== 3 =====================
h1("3. 货物进出口管理 — Импорт и экспорт товаров")
point("基本原则：自由进出口", "zìyóu jìnchūkǒu", "Базовый принцип — свободный импорт/экспорт товаров и технологий.")
point("3类货物", "sān lèi huòwù", "Товары делятся на 3 группы:")
point("自由、限制、禁止", "zìyóu, xiànzhì, jìnzhǐ", "свободные, ограниченные, запрещённые.")
point("3大手段", "sān dà shǒuduàn", "3 главных инструмента:")
point("许可证、配额、自动许可", "xǔkězhèng, pèi'é, zìdòng xǔkě",
      "лицензия, квота, автоматическое лицензирование.")
point("无证不准进出口", "wú zhèng bù zhǔn jìnchūkǒu",
      "Лицензия: без неё запрещён ввоз/вывоз ограниченных товаров.")
point("一关一证、一批一证", "yī guān yī zhèng, yī pī yī zhèng",
      "Лицензия: одна таможня — одна лицензия; одна партия — одна лицензия.")
point("关税配额：超额加税", "guānshuì pèi'é",
      "Тарифная квота: внутри квоты — низкий налог, сверх квоты — высокий (в основном для сельхозтоваров).")
point("自动许可：不能拒绝", "zìdòng xǔkě bùnéng jùjué",
      "Автолицензия — для СВОБОДНЫХ товаров, только для учёта; орган НЕ может отказать.")
note("Разница: лицензия = ограниченный товар, без неё нельзя. Автолицензия = свободный товар, отказать нельзя.")

# ===================== 4 =====================
h1("4. 主要环节 — Главные звенья (4 темы)")

p = doc.add_paragraph(); _set(p.add_run("A) 商品检验 — инспекция товаров"), CN_FONT, 12.5, bold=True, color=C_H1)
point("商检有5个原则", "5 ge yuánzé", "5 принципов инспекции:")
point("人、动植物、环境、防欺诈、国家安全", "",
      "здоровье людей; жизнь животных/растений; экология; борьба с обманом; нацбезопасность.")
point("3种检验", "3 zhǒng jiǎnyàn", "3 вида проверки: 法定 (обязательная), 免验 (освобождение), 抽查 (выборочная).")

p = doc.add_paragraph(); _set(p.add_run("B) 海关 — таможня"), CN_FONT, 12.5, bold=True, color=C_H1)
point("集中、统一、垂直", "jízhōng, tǒngyī, chuízhí",
      "Таможня: централизованная, единая, вертикальная; подчиняется только 海关总署.")
point("4种关税税率", "4 zhǒng shuìlǜ",
      "4 ставки пошлины: 最惠国 (наиб. благоприятств.), 协定 (договорная), 特惠 (преференц.), 普通 (обычная).")
point("完税价格 = 征税基础", "wánshuì jiàgé",
      "Таможенная стоимость — основа налога, считается по цене сделки (成交价格).")
point("查辑走私", "chájí zǒusī", "Борьба с контрабандой — одна из главных функций таможни (есть 缉私局).")

p = doc.add_paragraph(); _set(p.add_run("C) 原产地 — происхождение товара"), CN_FONT, 12.5, bold=True, color=C_H1)
point("2类规则", "2 lèi guīzé", "2 типа правил:")
point("非优惠（所有国）、优惠（协定国）", "fēiyōuhuì / yōuhuì",
      "непреференциальные (для всех стран) и преференциальные (только для стран-партнёров по соглашению).")
point("2个标准", "2 ge biāozhǔn", "2 критерия происхождения:")
point("完全获得、实质性改变", "wánquán huòdé / shízhìxìng gǎibiàn",
      "полностью получено в одной стране; или существенно переработано (тогда — страна последней переработки).")

p = doc.add_paragraph(); _set(p.add_run("D) 外汇 — валюта"), CN_FONT, 12.5, bold=True, color=C_H1)
point("管理外汇收支、汇率、市场", "wàihuì shōuzhī, huìlǜ, shìchǎng",
      "Государство контролирует валютные доходы/расходы, курс, валютный рынок.")
point("经常项目：不限制", "jīngcháng xiàngmù bù xiànzhì",
      "Текущие операции НЕ ограничиваются, но сделка должна быть真实、合法 (настоящей и законной).")
note("В этом разделе главное: 5 принципов商检, 4 ставки关税, 2 типа правил происхождения.")

# ===================== итоговая шпаргалка =====================
h1("★ Шпаргалка цифр (выучи это)")
make_table(
    ["数字 / число", "содержание"],
    [
        ["4 признака", "统一、速效、强制、纵向"],
        ["3 объекта", "经营者、货物、配套环节"],
        ["8 методов", "登记、国营贸易、许可证、配额、商检、原产地、海关、外汇"],
        ["3 недостатка рынка", "自发性、盲目性、滞后性"],
        ["许可制→备案登记制", "от разрешения к регистрации (право на ВЭД)"],
        ["8 импортных гостоваров", "粮食、植物油、糖、烟草、原油、成品油、化肥、棉花"],
        ["3 手段", "许可证、配额、自动许可"],
        ["两个一", "一关一证、一批一证"],
        ["5 принципов商检", "人、动植物、环境、防欺诈、国家安全"],
        ["3 检验", "法定、免验、抽查"],
        ["4 ставки关税", "最惠国、协定、特惠、普通"],
        ["2 правила原产地", "非优惠、优惠"],
        ["2 критерия", "完全获得、实质性改变"],
    ],
    [4.0, 12.0],
)

# ===================== мини-вопросы =====================
h1("★ Проверь себя (мини-тест)")
mini = [
    ("行政管理有哪4个特点？", "统一、速效、强制、纵向。"),
    ("市场有哪3个缺点？", "自发性、盲目性、滞后性。"),
    ("货物经营权现在是什么制度？", "备案登记制 (раньше — 许可制)。"),
    ("进口国营贸易有几种货物？", "8种：粮食、植物油、糖、烟草、原油、成品油、化肥、棉花。"),
    ("货物进出口3大手段？", "许可证、配额、自动许可。"),
    ("商检5原则？", "人、动植物、环境、防欺诈、国家安全。"),
    ("关税4种税率？", "最惠国、协定、特惠、普通。"),
    ("原产地2个标准？", "完全获得、实质性改变。"),
]
for i, (q, a) in enumerate(mini, 1):
    pq = doc.add_paragraph()
    _set(pq.add_run(f"{i}. {q}"), CN_FONT, 11.5, bold=True, color=RGBColor(0x1F,0x49,0x7D))
    pa = doc.add_paragraph(); pa.paragraph_format.left_indent = Cm(0.6)
    _set(pa.add_run("答：" + a), CN_FONT, 11)

out = "第六章 简明版 - 外国学生速记.docx"
doc.save(out)
print("Saved:", out, "| paras:", len(doc.paragraphs), "| tables:", len(doc.tables))
