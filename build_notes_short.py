# -*- coding: utf-8 -*-
"""
КРАТКАЯ (для лёгкого запоминания) версия конспекта
«第二章 对外贸易发展战略» — Стратегия развития внешней торговли Китая.

Принцип: только суть — короткие формулы, сводная таблица по пятилеткам,
мини-пояснения на русском (1-2 предложения) и компактный словарь.
"""

from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

C_TITLE = RGBColor(0x8B, 0x00, 0x00)
C_H2    = RGBColor(0xB0, 0x30, 0x10)
C_RU    = RGBColor(0x1A, 0x4D, 0x2E)
C_KEY   = RGBColor(0x8B, 0x00, 0x00)
C_GREY  = RGBColor(0x55, 0x55, 0x55)

CJK_FONT = "Microsoft YaHei"
LAT_FONT = "Calibri"

doc = Document()
style = doc.styles["Normal"]
style.font.name = LAT_FONT
style.font.size = Pt(11)
style._element.rPr.rFonts.set(qn("w:eastAsia"), CJK_FONT)
for s in doc.sections:
    s.top_margin = Cm(1.8); s.bottom_margin = Cm(1.8)
    s.left_margin = Cm(2.0); s.right_margin = Cm(2.0)


def _font(run, cjk=CJK_FONT, lat=LAT_FONT):
    run.font.name = lat
    rpr = run._element.get_or_add_rPr()
    rf = rpr.find(qn("w:rFonts"))
    if rf is None:
        rf = OxmlElement("w:rFonts"); rpr.append(rf)
    rf.set(qn("w:eastAsia"), cjk); rf.set(qn("w:ascii"), lat); rf.set(qn("w:hAnsi"), lat)


def shade(p, hexc):
    pPr = p._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), hexc)
    pPr.append(shd)


def lborder(p, hexc, size=20):
    pPr = p._p.get_or_add_pPr()
    pbdr = OxmlElement("w:pBdr"); left = OxmlElement("w:left")
    left.set(qn("w:val"), "single"); left.set(qn("w:sz"), str(size))
    left.set(qn("w:space"), "8"); left.set(qn("w:color"), hexc)
    pbdr.append(left); pPr.append(pbdr)


def run(p, text, size=11, bold=False, color=None, italic=False):
    r = p.add_run(text); r.font.size = Pt(size); r.bold = bold; r.italic = italic
    if color is not None:
        r.font.color.rgb = color
    _font(r); return r


def h1(cn, ru=None):
    p = doc.add_paragraph(); p.paragraph_format.space_before = Pt(12)
    run(p, cn, 16, True, C_TITLE)
    if ru:
        run(p, "  · " + ru, 10, False, C_GREY, italic=True)
    pPr = p._p.get_or_add_pPr(); pbdr = OxmlElement("w:pBdr"); b = OxmlElement("w:bottom")
    b.set(qn("w:val"), "single"); b.set(qn("w:sz"), "10"); b.set(qn("w:space"), "3"); b.set(qn("w:color"), "8B0000")
    pbdr.append(b); pPr.append(pbdr)


def h2(text):
    p = doc.add_paragraph(); p.paragraph_format.space_before = Pt(8); p.paragraph_format.space_after = Pt(2)
    run(p, text, 12.5, True, C_H2)


def b(text, lead=None):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.left_indent = Cm(0.55); p.paragraph_format.space_after = Pt(1)
    if lead:
        run(p, lead, 11, True, C_KEY)
    run(p, text, 11)


def ru(text):
    p = doc.add_paragraph(); p.paragraph_format.space_before = Pt(3); p.paragraph_format.space_after = Pt(7)
    shade(p, "EAF3EC"); lborder(p, "1A4D2E")
    run(p, "💡 ", 10, True, C_RU)
    run(p, text, 10, color=RGBColor(0x20, 0x20, 0x20))


# ---------------- Титул ----------------
t = doc.add_paragraph(); t.alignment = WD_ALIGN_PARAGRAPH.CENTER; t.paragraph_format.space_before = Pt(20)
run(t, "对外贸易发展战略", 26, True, C_TITLE)
s = doc.add_paragraph(); s.alignment = WD_ALIGN_PARAGRAPH.CENTER
run(s, "精简版 · 速记提纲 (краткий конспект для запоминания)", 12, True, RGBColor(0x1F, 0x3A, 0x5F))
s2 = doc.add_paragraph(); s2.alignment = WD_ALIGN_PARAGRAPH.CENTER
run(s2, "Стратегия развития внешней торговли Китая", 10.5, italic=True, color=C_GREY)
doc.add_paragraph()

# ---------------- Рамка ----------------
h1("六大战略", "6 стратегий — общая рамка")
b("出口商品战略 — экспортные товары (что вывозить)", "① ")
b("以质取胜 — победа за счёт качества", "② ")
b("科技兴贸 — торговля через науку и технику", "③ ")
b("出口市场多元化 — диверсификация рынков", "④ ")
b("进口商品战略 — импортные товары (что ввозить)", "⑤ ")
b("自由贸易区(FTA) — зоны свободной торговли / RCEP", "⑥ ")
ru("Логика всей главы: переход ОТ дешёвого труда и сырья (比较优势) "
   "К технологиям, брендам и качеству (竞争优势). Схема ответа: 背景→依据→内容→成效.")

# ---------------- Таблица по пятилеткам ----------------
h1("① 出口商品战略：按五年计划速记", "Эволюция экспорта — таблица")

rows = [
    ("六五", "1981-85", "发挥4优势：资源·传统技艺·劳动力·工业基础", "Опора на сырьё и труд"),
    ("七五", "1986-90", "“两个转变”：初级品→制成品；粗加工→精加工", "2 перехода (сырьё→изделия)"),
    ("八五", "1991-95", "第二个转变；机电产品成最大宗出口品", "Электромеханика — №1"),
    ("九五", "1996-2000", "提出“以质取胜”；粗放型→集约型", "Рождение «качество»"),
    ("十五", "2001-05", "继续“以质取胜”+“科技兴贸”", "Качество + технологии"),
    ("十一五", "2006-10", "“三自”(品牌·知识产权·营销)；控“两高一资”", "Свой бренд, контроль грязного"),
    ("十二五", "2011-15", "新优势：技术·品牌·质量·服务", "Подъём по цепочке стоимости"),
    ("十三五", "2016-20", "“优进优出”；跨境电商等新业态", "Качественно ввоз/вывоз"),
    ("十四五", "2021-25", "高质量发展·“三个变革”·数字贸易", "Цифра + высокое качество"),
]
tbl = doc.add_table(rows=1, cols=4); tbl.alignment = WD_TABLE_ALIGNMENT.CENTER; tbl.style = "Light Grid Accent 1"
for c, txt in zip(tbl.rows[0].cells, ["计划", "时期", "核心战略 / 口号", "Суть (RU)"]):
    c.paragraphs[0].clear(); r = c.paragraphs[0].add_run(txt)
    r.bold = True; r.font.size = Pt(10.5); r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF); _font(r)
    shd = OxmlElement("w:shd"); shd.set(qn("w:val"), "clear"); shd.set(qn("w:fill"), "8B0000")
    c._tc.get_or_add_tcPr().append(shd)
for plan, period, core, su in rows:
    cells = tbl.add_row().cells
    for cell, text, bold in ((cells[0], plan, True), (cells[1], period, False), (cells[2], core, False), (cells[3], su, False)):
        cell.paragraphs[0].clear(); r = cell.paragraphs[0].add_run(text)
        r.font.size = Pt(9.5); r.bold = bold; _font(r)
for rw in tbl.rows:
    rw.cells[0].width = Cm(1.6); rw.cells[1].width = Cm(2.0); rw.cells[2].width = Cm(8.2); rw.cells[3].width = Cm(4.8)
ru("Запомни лестницу: сначала 4 преимущества → «两个转变» → электромеханика лидер → «以质取胜» "
   "→ +科技兴贸 → «三自» → «优进优出» → цифровая торговля. Самый частый тест — соотнести пятилетку и лозунг.")

h2("重点对比：大进大出 → 优进优出")
b("大进大出：进口原料→国内加工→出口；中国只赚“打工者利润”")
b("优进：有选择进口紧缺的先进技术、关键设备、零部件")
b("优出：出口高附加值产品 + 推动“全产业链出口”(产品+技术+服务)")
ru("Главное противопоставление. Раньше Китай — «сборочный цех» (прибыль у иностранцев). "
   "优进优出 — ввозить и вывозить КАЧЕСТВЕННО, экспортируя всю цепочку, а не только сборку.")

doc.add_page_break()

# ---------------- 2 以质取胜 ----------------
h1("② 以质取胜", "Победа за счёт качества (с 九五)")
b("内涵(3)：提质量信誉 + 优化结构(高附加值) + 创名牌")
b("措施(5)：质量立法 · 提科技含量 · 接轨国际标准 · 名牌战略 · 全面质量管理")
ru("Качество = решающий фактор конкуренции. Структура ответа: «内涵(3)+措施(5)».")

# ---------------- 3 科技兴贸 ----------------
h1("③ 科技兴贸", "Торговля через науку и технику (十五)")
b("推动高新技术产品出口，培育竞争力强的企业")
b("用高新技术改造传统出口产业，提高其技术含量和附加值")
ru("Две стороны: РАЗВИВАТЬ новое (хайтек) + МОДЕРНИЗИРОВАТЬ старое. Идёт в паре с «以质取胜».")

# ---------------- 4 多元化 ----------------
h1("④ 出口市场多元化", "Диверсификация рынков")
b("四大传统市场：港澳 · 日本 · 美国 · 欧盟（过度集中→有风险）")
h2("必要性(4)")
b("减少贸易摩擦、规避风险")
b("出口持续稳定发展（避反倾销）")
b("争取有利贸易条件（防买方垄断压价）")
b("提升在国际分工中的地位（避免被锁在低端）")
h2("四类市场对策")
b("发达国家市场——深度开发(资金技术来源)，扩大水平分工")
b("亚洲市场——稳定扩大(东盟·韩国·香港转口)")
b("非洲拉美——开拓(适销对路+综合方式)")
b("中亚东欧——扩大(资源+相邻+互补)")
ru("4 рынка = 4 глагола: развитые «углублять», Азия «стабилизировать», Африка/ЛатАм «осваивать», "
   "ЦентрАзия/ВостЕвропа «расширять».")

# ---------------- 5 进口商品战略 ----------------
h1("⑤ 进口商品战略", "Импортные товары — что ввозить")
b("引进先进技术和关键设备（四个“确保”：电子信息·能源交通·传统产业改造·农业现代化）")
b("确保重要资源进口（钢材·橡胶·铜·金刚石·白金…）")
b("重视加工贸易物资进口（利用国外资源+国内劳动力，创汇、增就业）")
b("扩大生活必需品与一般消费品进口（如粮食→保国计民生）")
ru("Импорт опирается на цели нархозяйства. 4 уровня: технологии → ресурсы → "
   "сырьё для давальческой торговли → потребительские товары.")

# ---------------- 6 FTA ----------------
h1("⑥ 自由贸易区(FTA) / RCEP", "Зоны свободной торговли")
b("FTA：≥2国签协定，在WTO最惠国基础上进一步开放，分阶段取消关税与非关税壁垒")
b("一体化5级(自由化递增)：优惠贸易安排→自贸区→关税同盟→共同市场→经济同盟")
b("FTA是WTO的“例外”：成员间互给优惠，不必给其他成员；是补充也是深度开放")
h2("RCEP 必背")
b("全球最大自贸区；2020签署、2022生效")
b("15国 = 东盟10国 + 中日韩澳新（注意：印度未加入）")
b("覆盖≈全球30%人口·30%GDP·28%贸易；核心=区域内90%以上货物零关税")
ru("FTA — исключение из режима наибольшего благоприятствования (MFN). "
   "RCEP — крупнейшая ЗСТ: 15 стран (АСЕАН+КНР+Япония+Корея+Австралия+Н.Зеландия), Индии НЕТ. "
   "Подписан 2020, в силе с 2022, ≈30% мира.")

doc.add_page_break()

# ---------------- Мини-словарь ----------------
h1("速记词汇", "Мини-словарь ключевых терминов")
GL = [
    ("比较优势", "сравнительное преимущество"),
    ("竞争优势", "конкурентное преимущество"),
    ("初级产品 / 工业制成品", "сырьё / промышл. готовые изделия"),
    ("粗加工 / 精加工", "грубая / глубокая переработка"),
    ("两个转变", "«два перехода»"),
    ("机电产品", "электромеханическая продукция"),
    ("附加值", "добавленная стоимость"),
    ("劳动密集型 / 技术密集型", "трудоёмкий / технологоёмкий"),
    ("粗放型 / 集约型", "экстенсивный / интенсивный рост"),
    ("以质取胜", "победа за счёт качества"),
    ("名牌战略", "стратегия брендов"),
    ("科技兴贸", "торговля через науку и технику"),
    ("加工贸易", "давальческая (толлинговая) торговля"),
    ("三自(品牌·知识产权·营销)", "свой бренд / ИС / сбыт"),
    ("两高一资", "энергоёмкие, грязные, сырьевые товары"),
    ("优进优出 / 大进大出", "качественно / просто много ввоз-вывоз"),
    ("全产业链出口", "экспорт всей производств. цепочки"),
    ("跨境电子商务", "трансгранич. электронная торговля"),
    ("数字贸易", "цифровая торговля"),
    ("市场多元化", "диверсификация рынков"),
    ("贸易摩擦 / 反倾销", "торговые трения / антидемпинг"),
    ("买方垄断", "монополия покупателя"),
    ("国际分工", "международное разделение труда"),
    ("垂直 / 水平分工", "вертикальное / горизонтальное"),
    ("转口市场", "реэкспортный (транзитный) рынок"),
    ("贸易逆差", "дефицит (отриц. сальдо) торговли"),
    ("互补性", "взаимодополняемость"),
    ("自由贸易区(FTA)", "зона свободной торговли"),
    ("最惠国待遇", "режим наибольш. благоприятствования"),
    ("关税 / 非关税壁垒", "пошлина / нетарифные барьеры"),
    ("关税同盟 / 共同市场", "таможенный союз / общий рынок"),
    ("RCEP", "Всестор. региональное эконом. партнёрство"),
    ("零关税", "нулевая пошлина"),
    ("中间品", "промежуточные товары (компоненты)"),
]
gt = doc.add_table(rows=1, cols=2); gt.alignment = WD_TABLE_ALIGNMENT.CENTER; gt.style = "Light Grid Accent 1"
for c, txt in zip(gt.rows[0].cells, ["中文", "Русский"]):
    c.paragraphs[0].clear(); r = c.paragraphs[0].add_run(txt)
    r.bold = True; r.font.size = Pt(10.5); r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF); _font(r)
    shd = OxmlElement("w:shd"); shd.set(qn("w:val"), "clear"); shd.set(qn("w:fill"), "8B0000")
    c._tc.get_or_add_tcPr().append(shd)
for zh, rr in GL:
    cells = gt.add_row().cells
    cells[0].paragraphs[0].clear(); a = cells[0].paragraphs[0].add_run(zh); a.bold = True; a.font.size = Pt(10); _font(a)
    cells[1].paragraphs[0].clear(); c2 = cells[1].paragraphs[0].add_run(rr); c2.font.size = Pt(10); _font(c2)
for rw in gt.rows:
    rw.cells[0].width = Cm(6.0); rw.cells[1].width = Cm(9.0)

doc.add_paragraph()
f = doc.add_paragraph(); f.alignment = WD_ALIGN_PARAGRAPH.CENTER
run(f, "— 加油！Удачи! —", 12, True, C_TITLE)

out = "对外贸易发展战略_精简速记版.docx"
doc.save(out)
print("Сохранён:", out, "| параграфов:", len(doc.paragraphs), "| терминов:", len(GL))
